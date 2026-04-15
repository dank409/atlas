from __future__ import annotations

import json
import re
import sqlite3
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "raw"
OUTPUT_DIR = ROOT / "public" / "data" / "infrastructure"

TARGET_COUNTRIES = {
    "Saudi Arabia": "SA",
    "Iran": "IR",
    "Iraq": "IQ",
    "Kuwait": "KW",
    "United Arab Emirates": "AE",
    "Qatar": "QA",
    "Oman": "OM",
    "Bahrain": "BH",
    "Yemen": "YE",
    "Egypt": "EG",
    "Syria": "SY",
    "Jordan": "JO",
    "Turkey": "TR",
    "Pakistan": "PK",
    "India": "IN",
}

COMPANY_ALIASES = {
    "saudi arabian oil": "saudi-aramco",
    "saudi arabian oil company": "saudi-aramco",
    "saudi aramco": "saudi-aramco",
    "aramco": "saudi-aramco",
    "abu dhabi national oil co": "adnoc",
    "abu dhabi national oil company": "adnoc",
    "adnoc": "adnoc",
    "kuwait petroleum corp": "kuwait-petroleum",
    "kuwait petroleum corporation": "kuwait-petroleum",
    "kuwait national petroleum co": "kuwait-petroleum",
    "kpc": "kuwait-petroleum",
}

TERMINAL_STATUS_PRIORITY = {
    "operating": 0,
    "construction": 1,
    "proposed": 2,
    "idled": 3,
    "mothballed": 4,
    "shelved": 5,
    "retired": 6,
    "cancelled": 7,
    "unknown": 8,
}


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    refineries = build_refineries()
    fields = build_fields()
    terminals = build_terminals()
    pipelines = build_pipelines()

    write_feature_collection(OUTPUT_DIR / "refineries.geojson", refineries)
    write_feature_collection(OUTPUT_DIR / "fields.geojson", fields)
    write_feature_collection(OUTPUT_DIR / "terminals.geojson", terminals)
    write_feature_collection(OUTPUT_DIR / "pipelines.geojson", pipelines)

    print(f"Wrote {len(refineries['features'])} refineries")
    print(f"Wrote {len(fields['features'])} fields")
    print(f"Wrote {len(terminals['features'])} terminals")
    print(f"Wrote {len(pipelines['features'])} pipelines")


def build_refineries() -> dict[str, Any]:
    database_path = RAW_DIR / "OGIM" / "OGIM_v2.7.gpkg"
    supplemental_path = RAW_DIR / "reference" / "refinery_seed_supplement.json"
    connection = sqlite3.connect(database_path)
    connection.row_factory = sqlite3.Row
    rows = connection.execute(
        """
        SELECT
            COUNTRY,
            FAC_NAME,
            FAC_TYPE,
            FAC_STATUS,
            OGIM_STATUS,
            OPERATOR,
            LIQ_CAPACITY_BPD,
            LIQ_THROUGHPUT_BPD,
            LATITUDE,
            LONGITUDE
        FROM Crude_Oil_Refineries
        """
    ).fetchall()
    connection.close()

    features: list[dict[str, Any]] = []
    used_feature_ids: set[str] = set()

    for row in rows:
        country_name = normalize_ogim_country(clean_string(row["COUNTRY"]))
        if country_name not in TARGET_COUNTRIES:
            continue

        lat = positive_or_none(row["LATITUDE"], allow_zero=False)
        lon = to_float(row["LONGITUDE"])
        if lat is None or lon is None:
            continue

        capacity_bpd = first_positive_number(row["LIQ_CAPACITY_BPD"], row["LIQ_THROUGHPUT_BPD"])
        operator = clean_entity_name(row["OPERATOR"]) or "Unknown"
        country_code = TARGET_COUNTRIES[country_name]
        name = format_refinery_name(clean_string(row["FAC_NAME"]) or "Unknown refinery")
        status = normalize_refinery_status(row["OGIM_STATUS"], row["FAC_STATUS"])
        feature_id = build_unique_refinery_id(
            used_feature_ids,
            country_code,
            name,
            operator,
            lat,
            lon,
        )

        properties: dict[str, Any] = {
            "id": feature_id,
            "name": name,
            "country": country_code,
            "countryName": country_name,
            "owner": operator,
            "operator": operator,
            "status": status,
            "capacityUnit": "bpd",
            "type": "refinery",
            "lat": lat,
            "lon": lon,
            "firmsRadius_km": refinery_firms_radius(capacity_bpd),
            "linkedContext": build_linked_context(country_code, operator),
        }

        if capacity_bpd is not None:
            properties["capacity_bpd"] = round(capacity_bpd)

        fac_type = clean_string(row["FAC_TYPE"])
        if fac_type and fac_type != "N/A":
            properties["notes"] = fac_type

        features.append(
            {
                "type": "Feature",
                "id": feature_id,
                "geometry": {
                    "type": "Point",
                    "coordinates": [lon, lat],
                },
                "properties": properties,
            }
        )

    if supplemental_path.exists():
        supplemental_payload = json.loads(supplemental_path.read_text(encoding="utf-8"))
        for feature in supplemental_payload.get("features", []):
            properties = feature.get("properties") or {}
            country_name = clean_string(properties.get("countryName"))
            if country_name not in TARGET_COUNTRIES:
                continue
            if refinery_exists(features, properties):
                continue

            lat = to_float(properties.get("lat"))
            lon = to_float(properties.get("lon"))
            if lat is None or lon is None:
                continue

            capacity_bpd = positive_or_none(properties.get("capacity_bpd"), allow_zero=False)
            operator = clean_entity_name(properties.get("operator")) or "Unknown"
            owner = clean_entity_name(properties.get("owner")) or operator
            country_code = TARGET_COUNTRIES[country_name]
            name = clean_string(properties.get("name")) or "Unknown refinery"
            feature_id = build_unique_refinery_id(
                used_feature_ids,
                country_code,
                name,
                operator,
                lat,
                lon,
                clean_string(properties.get("id")),
            )

            supplemental_properties: dict[str, Any] = {
                "id": feature_id,
                "name": name,
                "country": country_code,
                "countryName": country_name,
                "owner": owner,
                "operator": operator,
                "status": "unknown",
                "capacityUnit": "bpd",
                "type": "refinery",
                "lat": lat,
                "lon": lon,
                "firmsRadius_km": refinery_firms_radius(capacity_bpd),
                "linkedContext": build_linked_context(country_code, operator, owner),
                "notes": "Supplemental refinery site retained because OGIM v2.7 does not include this facility."
            }

            if capacity_bpd is not None:
                supplemental_properties["capacity_bpd"] = round(capacity_bpd)

            features.append(
                {
                    "type": "Feature",
                    "id": feature_id,
                    "geometry": {
                        "type": "Point",
                        "coordinates": [lon, lat],
                    },
                    "properties": supplemental_properties,
                }
            )

    features = sorted(features, key=lambda item: (item["properties"]["countryName"], item["properties"]["name"]))

    return {
        "type": "FeatureCollection",
        "_metadata": {
            "source": "EDF Oil and Gas Infrastructure Mapping (OGIM) v2.7 + curated refinery supplement",
            "downloadDate": "2025-03-28",
            "recordCount": len(features),
            "license": "CC BY 4.0",
            "sourceUrl": "https://zenodo.org/records/15103476",
            "notes": "OGIM refinery status is largely unavailable in the regional slice, so unknown values are preserved instead of inferred. Curated supplement rows are only used when named Gulf sites are absent from OGIM."
        },
        "features": features,
    }


def build_fields() -> dict[str, Any]:
    workbook_path = RAW_DIR / "Global-Oil-and-Gas-Extraction-Tracker-March-2026.xlsx"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    main_sheet = workbook["Field-level main data"]
    headers = next(main_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    index = header_index(headers)
    production_by_unit = load_field_production(workbook)

    features: list[dict[str, Any]] = []

    for row in main_sheet.iter_rows(min_row=2, values_only=True):
        country_name = clean_string(row[index["Country/Area"]])
        if country_name not in TARGET_COUNTRIES:
            continue

        lat = to_float(row[index["Latitude"]])
        lon = to_float(row[index["Longitude"]])

        if lat is None or lon is None:
            continue

        unit_id = clean_string(row[index["Unit ID"]]) or slugify(clean_string(row[index["Unit Name"]]) or "field")
        operator = clean_entity_name(row[index["Operator"]])
        country_code = TARGET_COUNTRIES[country_name]

        properties = {
            "id": f"field-{unit_id.lower()}",
            "name": clean_string(row[index["Unit Name"]]) or unit_id,
            "country": country_code,
            "countryName": country_name,
            "status": normalize_status(row[index["Status"]]),
            "operator": operator,
            "basin": clean_string(row[index["Basin"]]),
            "fieldType": normalize_field_type(row[index["Fuel type"]]),
            "lat": lat,
            "lon": lon,
            "linkedContext": build_linked_context(country_code, operator),
        }

        production_bpd = production_by_unit.get(unit_id)
        if production_bpd is not None:
            properties["production_bpd"] = round(production_bpd)

        features.append(
            {
                "type": "Feature",
                "id": properties["id"],
                "geometry": {
                    "type": "Point",
                    "coordinates": [lon, lat],
                },
                "properties": properties,
            }
        )

    return {
        "type": "FeatureCollection",
        "_metadata": {
            "source": "Global Energy Monitor - Global Oil and Gas Extraction Tracker",
            "downloadDate": "2026-03",
            "recordCount": len(features),
            "license": "CC BY 4.0",
            "sourceUrl": "https://globalenergymonitor.org/projects/global-oil-gas-extraction-tracker/download-data/",
        },
        "features": features,
    }


def build_terminals() -> dict[str, Any]:
    workbook_path = RAW_DIR / "GEM-GGIT-LNG-Teminals-2025-09.xlsx"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["LNG Terminals"]
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    index = header_index(headers)

    grouped_rows: dict[str, list[tuple[Any, ...]]] = defaultdict(list)
    for row in sheet.iter_rows(min_row=2, values_only=True):
        country_name = clean_string(row[index["Country/Area"]])
        if country_name not in TARGET_COUNTRIES:
            continue

        project_id = clean_string(row[index["ProjectID"]])
        if not project_id:
            continue

        grouped_rows[project_id].append(row)

    features: list[dict[str, Any]] = []

    for project_id, rows in grouped_rows.items():
        representative = pick_terminal_representative(rows, index)
        if representative is None:
            continue

        country_name = clean_string(representative[index["Country/Area"]])
        country_code = TARGET_COUNTRIES[country_name]
        lat = to_float(representative[index["Latitude"]])
        lon = to_float(representative[index["Longitude"]])
        if lat is None or lon is None:
            continue

        owner = first_non_empty(clean_entity_name(row[index["Owner"]]) for row in rows)
        operator = first_non_empty(clean_entity_name(row[index["Operator"]]) for row in rows)
        facility_type = first_non_empty(clean_string(row[index["FacilityType"]]) for row in rows)

        properties: dict[str, Any] = {
            "id": f"terminal-{project_id.lower()}",
            "name": clean_string(representative[index["TerminalName"]]) or project_id,
            "country": country_code,
            "countryName": country_name,
            "status": choose_terminal_status(rows, index),
            "lat": lat,
            "lon": lon,
            "owner": owner,
            "operator": operator,
            "terminalType": normalize_terminal_type(facility_type),
            "terminalRole": normalize_terminal_role(facility_type),
            "linkedContext": build_linked_context(country_code, operator, owner),
        }

        capacity_mtpa = derive_terminal_capacity_mtpa(rows, index)
        if capacity_mtpa is not None:
            properties["capacity_mtpa"] = round(capacity_mtpa, 2)

        features.append(
            {
                "type": "Feature",
                "id": properties["id"],
                "geometry": {
                    "type": "Point",
                    "coordinates": [lon, lat],
                },
                "properties": properties,
            }
        )

    return {
        "type": "FeatureCollection",
        "_metadata": {
            "source": "Global Energy Monitor - Global Gas Infrastructure Tracker (LNG terminals)",
            "downloadDate": "2025-09",
            "recordCount": len(features),
            "license": "CC BY 4.0",
            "sourceUrl": "https://globalenergymonitor.org/projects/global-gas-infrastructure-tracker/download-data/",
        },
        "features": sorted(features, key=lambda feature: feature["properties"]["name"]),
    }


def build_pipelines() -> dict[str, Any]:
    sources = [
        (
            RAW_DIR / "GEM-GOIT-Oil-NGL-Pipelines-2025-03" / "GEM-GOIT-Oil-NGL-Pipelines-2025-03.geojson",
            "Countries",
            "Global Energy Monitor - Global Oil Infrastructure Tracker",
            "2025-03",
        ),
        (
            RAW_DIR / "GEM-GGIT-Gas-Pipelines-2025-11" / "GEM-GGIT-Gas-Pipelines-2025-11.geojson",
            "CountriesOrAreas",
            "Global Energy Monitor - Global Gas Infrastructure Tracker",
            "2025-11",
        ),
    ]

    features: list[dict[str, Any]] = []

    for path, country_field, source_name, download_date in sources:
        with path.open() as handle:
            payload = json.load(handle)

        for feature in payload.get("features", []):
            props = feature.get("properties") or {}
            countries = [
                country.strip()
                for country in str(props.get(country_field) or "").split(",")
                if country and country.strip() in TARGET_COUNTRIES
            ]

            if not countries:
                continue

            geometry = normalize_pipeline_geometry(feature.get("geometry"))
            if geometry is None:
                continue

            primary_country = countries[0]
            country_code = TARGET_COUNTRIES[primary_country]
            owner = clean_entity_name(props.get("Owner"))
            parent = clean_entity_name(props.get("Parent"))

            properties: dict[str, Any] = {
                "id": f"pipeline-{str(props.get('ProjectID') or slugify(props.get('PipelineName') or 'pipeline')).lower()}",
                "name": clean_string(props.get("PipelineName")) or clean_string(props.get("SegmentName")) or "Unknown pipeline",
                "country": country_code,
                "countryName": primary_country,
                "countries": [TARGET_COUNTRIES[country] for country in countries],
                "owner": owner or parent or "Unknown",
                "operator": None,
                "status": normalize_status(props.get("Status")),
                "hydrocarbon": normalize_hydrocarbon(props.get("Fuel")),
                "capacity_boed": to_float(props.get("CapacityBOEd")),
                "diameter_inches": parse_diameter_inches(props.get("Diameter")),
                "length_km": first_number(
                    props.get("LengthMergedKm"),
                    props.get("LengthKnownKm"),
                    props.get("LengthEstimateKm"),
                ),
                "startPoint": clean_string(props.get("StartLocation")),
                "endPoint": clean_string(props.get("EndLocation")),
                "segmentName": clean_string(props.get("SegmentName")),
                "linkedContext": build_linked_context(country_code, owner, parent),
            }

            properties = {key: value for key, value in properties.items() if value is not None}

            features.append(
                {
                    "type": "Feature",
                    "id": properties["id"],
                    "geometry": geometry,
                    "properties": properties,
                    "_source": {
                        "downloadDate": download_date,
                        "source": source_name,
                    },
                }
            )

    serialized_features = []
    for feature in sorted(features, key=lambda item: item["properties"]["name"]):
        source = feature.pop("_source")
        feature["properties"]["_ingestSource"] = source["source"]
        feature["properties"]["_ingestDownloadDate"] = source["downloadDate"]
        serialized_features.append(feature)

    return {
        "type": "FeatureCollection",
        "_metadata": {
            "source": "Global Energy Monitor - Global Oil Infrastructure Tracker / Global Gas Infrastructure Tracker",
            "downloadDate": "2025-11",
            "recordCount": len(serialized_features),
            "license": "CC BY 4.0",
            "sourceUrl": "https://globalenergymonitor.org/projects/global-oil-infrastructure-tracker/download-data/",
        },
        "features": serialized_features,
    }


def load_field_production(workbook) -> dict[str, float]:
    sheet = workbook["Field-level production data"]
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    index = header_index(headers)
    grouped_rows: dict[str, list[tuple[Any, ...]]] = defaultdict(list)

    for row in sheet.iter_rows(min_row=2, values_only=True):
        country_name = clean_string(row[index["Country/Area"]])
        unit_id = clean_string(row[index["Unit ID"]])
        if country_name not in TARGET_COUNTRIES or not unit_id:
            continue
        grouped_rows[unit_id].append(row)

    production_by_unit: dict[str, float] = {}

    for unit_id, rows in grouped_rows.items():
        years = [to_int(row[index["Data Year"]]) for row in rows if to_int(row[index["Data Year"]]) is not None]
        if not years:
            continue
        latest_year = max(years)
        values = [
            production_row_to_bpd(row, index)
            for row in rows
            if to_int(row[index["Data Year"]]) == latest_year
        ]
        total = sum(value for value in values if value is not None)
        if total > 0:
            production_by_unit[unit_id] = total

    return production_by_unit


def production_row_to_bpd(row: tuple[Any, ...], index: dict[str, int]) -> float | None:
    original_value = to_float(row[index["Quantity (original)"]])
    original_unit = clean_string(row[index["Units (original)"]])
    converted_value = to_float(row[index["Quantity (converted)"]])
    converted_unit = clean_string(row[index["Units (converted)"]])

    if original_value is not None and original_unit == "bbl/d":
        return original_value
    if original_value is not None and original_unit == "thousand bbl/d":
        return original_value * 1000
    if original_value is not None and original_unit == "boe/d":
        return original_value
    if converted_value is not None and converted_unit == "million bbl/y":
        return converted_value * 1_000_000 / 365
    if converted_value is not None and converted_unit == "million boe/y":
        return converted_value * 1_000_000 / 365

    return None


def pick_terminal_representative(rows: list[tuple[Any, ...]], index: dict[str, int]) -> tuple[Any, ...] | None:
    best_row = None
    best_score = None

    for row in rows:
        lat = to_float(row[index["Latitude"]])
        lon = to_float(row[index["Longitude"]])
        if lat is None or lon is None:
            continue

        status = normalize_status(row[index["Status"]])
        score = TERMINAL_STATUS_PRIORITY.get(status, TERMINAL_STATUS_PRIORITY["unknown"])

        if best_row is None or score < best_score:
            best_row = row
            best_score = score

    return best_row


def choose_terminal_status(rows: list[tuple[Any, ...]], index: dict[str, int]) -> str:
    statuses = [normalize_status(row[index["Status"]]) for row in rows]
    return min(statuses, key=lambda status: TERMINAL_STATUS_PRIORITY.get(status, TERMINAL_STATUS_PRIORITY["unknown"]))


def derive_terminal_capacity_mtpa(rows: list[tuple[Any, ...]], index: dict[str, int]) -> float | None:
    totals: list[float] = []
    unit_capacities: list[float] = []

    for row in rows:
        for field in ("TotImportLNGTerminalCapacityinMtpa", "TotExportLNGTerminalCapacityinMtpa"):
            value = to_float(row[index[field]])
            if value is not None:
                totals.append(value)

        capacity = to_float(row[index["CapacityinMtpa"]])
        if capacity is not None:
            unit_capacities.append(capacity)

    if totals:
        return max(totals)
    if unit_capacities:
        return sum(unit_capacities)

    return None


def normalize_pipeline_geometry(geometry: dict[str, Any] | None) -> dict[str, Any] | None:
    if not geometry:
        return None

    geometry_type = geometry.get("type")

    if geometry_type == "LineString" and geometry.get("coordinates"):
        return {
            "type": "LineString",
            "coordinates": geometry["coordinates"],
        }

    if geometry_type == "MultiLineString" and geometry.get("coordinates"):
        return {
            "type": "MultiLineString",
            "coordinates": geometry["coordinates"],
        }

    if geometry_type == "GeometryCollection":
        line_parts: list[list[list[float]]] = []
        for item in geometry.get("geometries", []):
            item_type = item.get("type")
            if item_type == "LineString" and item.get("coordinates"):
                line_parts.append(item["coordinates"])
            elif item_type == "MultiLineString" and item.get("coordinates"):
                line_parts.extend(item["coordinates"])

        if len(line_parts) == 1:
            return {
                "type": "LineString",
                "coordinates": line_parts[0],
            }
        if len(line_parts) > 1:
            return {
                "type": "MultiLineString",
                "coordinates": line_parts,
            }

    return None


def build_linked_context(country_code: str, *entities: str | None) -> dict[str, Any] | None:
    company_id = None
    for entity in entities:
        normalized = normalize_name(entity)
        if not normalized:
            continue

        if normalized in COMPANY_ALIASES:
            company_id = COMPANY_ALIASES[normalized]
            break

        for alias, candidate_company_id in COMPANY_ALIASES.items():
            if alias in normalized:
                company_id = candidate_company_id
                break

        if company_id:
            break

    context: dict[str, Any] = {
        "countryCode": country_code,
    }
    if company_id:
        context["companyId"] = company_id

    return context


def normalize_field_type(value: Any) -> str | None:
    normalized = normalize_name(value)
    if not normalized:
        return None
    if "gas" in normalized and "oil" in normalized:
        return "mixed"
    if "condensate" in normalized:
        return "condensate"
    if "gas" in normalized:
        return "gas"
    if "oil" in normalized:
        return "oil"
    return "mixed"


def normalize_terminal_type(value: str | None) -> str | None:
    if value == "export":
        return "export_terminal"
    if value == "import":
        return "import_terminal"
    return "lng_terminal"


def normalize_terminal_role(value: str | None) -> str | None:
    if value in {"export", "import"}:
        return value
    return None


def normalize_hydrocarbon(value: Any) -> str:
    normalized = normalize_name(value)
    if normalized == "gas":
        return "natural_gas"
    if normalized == "ngl":
        return "ngl"
    if normalized == "oil":
        return "crude_oil"
    if "oil" in normalized and "gas" in normalized:
        return "mixed"
    return "mixed"


def normalize_ogim_country(value: str | None) -> str | None:
    if not value:
        return None

    normalized = normalize_name(value)
    country_aliases = {
        "united arab emirates": "United Arab Emirates",
        "uae": "United Arab Emirates",
        "saudi arabia": "Saudi Arabia",
        "iran": "Iran",
        "iraq": "Iraq",
        "kuwait": "Kuwait",
        "qatar": "Qatar",
        "oman": "Oman",
        "bahrain": "Bahrain",
        "yemen": "Yemen",
        "egypt": "Egypt",
        "syria": "Syria",
        "jordan": "Jordan",
        "turkey": "Turkey",
        "pakistan": "Pakistan",
        "india": "India",
    }
    return country_aliases.get(normalized)


def normalize_refinery_status(*values: Any) -> str:
    for value in values:
        normalized = normalize_name(value)
        if not normalized or normalized == "n a":
            continue
        if "incident" in normalized or "outage" in normalized:
            return "unplanned_incident"
        if "maintenance" in normalized:
            return "planned_maintenance"
        if "construct" in normalized or "build" in normalized:
            return "construction"
        if "proposed" in normalized or "planned" in normalized or "startup" in normalized:
            return "planned_startup"
        if "clos" in normalized or "retired" in normalized or "decommission" in normalized:
            return "permanent_closure"
        if "operat" in normalized or "active" in normalized or "online" in normalized:
            return "operational"

    return "unknown"


def normalize_status(value: Any) -> str:
    cleaned = normalize_name(value)
    if not cleaned:
        return "unknown"
    return cleaned.replace("-", "_").replace(" ", "_")


def clean_entity_name(value: Any) -> str | None:
    cleaned = clean_string(value)
    if not cleaned or cleaned == "--":
        return None
    first = cleaned.split(";")[0].strip()
    first = re.sub(r"\s*\[[^\]]*\]", "", first).strip()
    return first or None


def clean_string(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).strip()
    return text or None


def normalize_name(value: Any) -> str:
    text = clean_string(value) or ""
    text = text.casefold()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def parse_diameter_inches(value: Any) -> float | None:
    cleaned = clean_string(value)
    if not cleaned:
        return None
    numbers = [float(match) for match in re.findall(r"\d+(?:\.\d+)?", cleaned)]
    if not numbers:
        return None
    return max(numbers)


def first_number(*values: Any) -> float | None:
    for value in values:
        numeric = to_float(value)
        if numeric is not None:
            return numeric
    return None


def first_positive_number(*values: Any) -> float | None:
    for value in values:
        numeric = positive_or_none(value, allow_zero=False)
        if numeric is not None:
            return numeric
    return None


def header_index(headers: tuple[Any, ...]) -> dict[str, int]:
    return {str(header): position for position, header in enumerate(headers)}


def to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if text in {"", "--"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def positive_or_none(value: Any, allow_zero: bool = False) -> float | None:
    numeric = to_float(value)
    if numeric is None:
        return None
    if allow_zero:
        return numeric if numeric >= 0 else None
    return numeric if numeric > 0 else None


def to_int(value: Any) -> int | None:
    numeric = to_float(value)
    if numeric is None:
        return None
    return int(numeric)


def slugify(value: Any) -> str:
    normalized = normalize_name(value)
    return normalized.replace(" ", "-") or "item"


def first_non_empty(values):
    for value in values:
        if value:
            return value
    return None


def refinery_exists(features: list[dict[str, Any]], candidate_properties: dict[str, Any]) -> bool:
    candidate_country = clean_string(candidate_properties.get("countryName")) or ""
    candidate_name = clean_string(candidate_properties.get("name")) or ""
    candidate_normalized = normalize_name(candidate_name)

    for feature in features:
        existing = feature["properties"]
        if existing["countryName"] != candidate_country:
            continue

        existing_normalized = normalize_name(existing["name"])
        if (
            existing_normalized == candidate_normalized
            or existing_normalized in candidate_normalized
            or candidate_normalized in existing_normalized
        ):
            return True

    return False


def build_unique_refinery_id(
    used_feature_ids: set[str],
    country_code: str,
    name: str,
    operator: str,
    lat: float,
    lon: float,
    preferred_id: str | None = None,
) -> str:
    candidates = [
        preferred_id,
        f"refinery-{country_code.lower()}-{slugify(name)}",
        f"refinery-{country_code.lower()}-{slugify(name)}-{slugify(operator)}",
        f"refinery-{country_code.lower()}-{slugify(name)}-{format_coordinate_suffix(lat, lon)}",
    ]

    for candidate in candidates:
        if not candidate or candidate in used_feature_ids:
            continue

        used_feature_ids.add(candidate)
        return candidate

    suffix = 2
    while True:
        candidate = f"refinery-{country_code.lower()}-{slugify(name)}-{suffix}"
        if candidate not in used_feature_ids:
            used_feature_ids.add(candidate)
            return candidate
        suffix += 1


def format_coordinate_suffix(lat: float, lon: float) -> str:
    return f"{abs(round(lat * 100)):04d}-{abs(round(lon * 100)):04d}"


def format_refinery_name(value: str) -> str:
    normalized = clean_string(value) or value
    titled = normalized.title()
    return titled if "refinery" in normalize_name(titled) else f"{titled} Refinery"


def refinery_firms_radius(capacity_bpd: float | None) -> int:
    if capacity_bpd is None:
        return 3
    if capacity_bpd >= 200_000:
        return 5
    if capacity_bpd >= 50_000:
        return 3
    return 2


def write_feature_collection(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
